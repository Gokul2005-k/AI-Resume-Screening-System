import io
import re
from pathlib import Path
from typing import Any

from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS, TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


SKILL_CATALOG = {
    "Programming": [
        "python",
        "java",
        "c++",
        "c",
        "javascript",
        "typescript",
        "sql",
        "html",
        "css",
    ],
    "Data": [
        "machine learning",
        "deep learning",
        "nlp",
        "data analysis",
        "data science",
        "pandas",
        "numpy",
        "tensorflow",
        "pytorch",
        "scikit-learn",
        "power bi",
        "tableau",
        "excel",
    ],
    "CloudAndDevOps": [
        "aws",
        "azure",
        "gcp",
        "docker",
        "kubernetes",
        "jenkins",
        "git",
        "github",
        "linux",
        "ci/cd",
    ],
    "Backend": [
        "django",
        "flask",
        "fastapi",
        "rest api",
        "api",
        "microservices",
        "mongodb",
        "mysql",
        "postgresql",
    ],
    "SoftSkills": [
        "communication",
        "leadership",
        "teamwork",
        "problem solving",
        "analytical",
        "time management",
    ],
}


ALL_SKILLS = sorted({skill for skills in SKILL_CATALOG.values() for skill in skills})

ACTION_VERBS = [
    "built",
    "developed",
    "designed",
    "implemented",
    "created",
    "optimized",
    "deployed",
    "automated",
    "analyzed",
    "improved",
    "led",
]

EXPECTED_SECTIONS = {
    "summary": ["summary", "profile", "objective", "about"],
    "skills": ["skills", "technical skills", "key skills"],
    "experience": ["experience", "internship", "work history", "employment"],
    "projects": ["project", "projects"],
    "education": ["education", "degree", "college", "university"],
}


def normalize_text(text: str) -> str:
    lowered = text.lower()
    cleaned = re.sub(r"[^a-z0-9+#./\s-]", " ", lowered)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned


def extract_skills(text: str) -> list[str]:
    normalized = normalize_text(text)
    matched = []
    for skill in ALL_SKILLS:
        pattern = rf"(?<!\w){re.escape(skill)}(?!\w)"
        if re.search(pattern, normalized):
            matched.append(skill)
    return sorted(set(matched))


def extract_keywords(text: str, top_n: int = 15) -> list[str]:
    normalized = normalize_text(text)
    words = [
        token
        for token in normalized.split()
        if token not in ENGLISH_STOP_WORDS and len(token) > 2
    ]
    frequencies: dict[str, int] = {}
    for word in words:
        frequencies[word] = frequencies.get(word, 0) + 1
    ranked = sorted(frequencies.items(), key=lambda item: (-item[1], item[0]))
    return [word for word, _ in ranked[:top_n]]


def analyze_resume_structure(resume_text: str) -> dict[str, Any]:
    normalized_resume = normalize_text(resume_text)
    word_count = len(normalized_resume.split())
    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", resume_text))
    has_phone = bool(re.search(r"(\+?\d[\d\s().-]{7,}\d)", resume_text))
    has_metrics = bool(re.search(r"\b\d+%|\b\d+\+|\b\d+\b", resume_text))
    has_action_verbs = any(verb in normalized_resume for verb in ACTION_VERBS)

    sections = {
        section: any(marker in normalized_resume for marker in markers)
        for section, markers in EXPECTED_SECTIONS.items()
    }

    checks = [
        {
            "label": "Contact details",
            "status": "Passed" if has_email or has_phone else "Needs work",
            "detail": (
                "Email or phone number is visible."
                if has_email or has_phone
                else "Add a clear email and phone number near the top."
            ),
        },
        {
            "label": "Skills section",
            "status": "Passed" if sections["skills"] else "Needs work",
            "detail": (
                "A skills section is available for quick scanning."
                if sections["skills"]
                else "Add a separate skills section with role-specific tools."
            ),
        },
        {
            "label": "Project or experience proof",
            "status": (
                "Passed"
                if sections["projects"] or sections["experience"]
                else "Needs work"
            ),
            "detail": (
                "The resume contains practical work signals."
                if sections["projects"] or sections["experience"]
                else "Add projects, internships, or practical work examples."
            ),
        },
        {
            "label": "Measured impact",
            "status": "Passed" if has_metrics else "Needs work",
            "detail": (
                "Numbers or measurable outcomes are present."
                if has_metrics
                else "Add metrics such as accuracy, users, duration, or improvement."
            ),
        },
        {
            "label": "Action-oriented wording",
            "status": "Passed" if has_action_verbs else "Needs work",
            "detail": (
                "Action verbs make the resume sound stronger."
                if has_action_verbs
                else "Start bullets with verbs like built, developed, or optimized."
            ),
        },
    ]

    passed_checks = sum(1 for check in checks if check["status"] == "Passed")
    readability_score = round((passed_checks / len(checks)) * 100)

    if word_count < 80:
        length_message = "Resume text looks too short for a complete ATS scan."
    elif word_count > 900:
        length_message = "Resume text may be too long. Keep it focused and role-specific."
    else:
        length_message = "Resume length looks reasonable for a first-pass scan."

    return {
        "word_count": word_count,
        "readability_score": readability_score,
        "length_message": length_message,
        "sections": sections,
        "checks": checks,
    }


def compute_similarity_score(resume_text: str, job_text: str) -> int:
    vectorizer = TfidfVectorizer(ngram_range=(1, 2), stop_words="english")
    tfidf_matrix = vectorizer.fit_transform([resume_text, job_text])
    similarity = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]
    return round(similarity * 100)


def generate_strengths(
    similarity_score: int, matched_skills: list[str], resume_text: str
) -> list[str]:
    strengths = []
    if similarity_score >= 70:
        strengths.append("Strong overall semantic alignment with the target role.")
    elif similarity_score >= 50:
        strengths.append("Moderate alignment with the job description.")
    else:
        strengths.append("The resume can be aligned more closely with the role language.")

    if len(matched_skills) >= 5:
        strengths.append("The resume already covers several core technical skills.")
    elif matched_skills:
        strengths.append("The resume contains some of the required skills.")
    else:
        strengths.append("Direct skill overlap is currently limited.")

    if re.search(r"\b\d+%|\b\d+\+|\b\d+\b", resume_text):
        strengths.append("The resume includes measurable or quantified achievements.")
    else:
        strengths.append("Adding metrics can make project and experience impact clearer.")

    return strengths


def generate_suggestions(
    missing_skills: list[str],
    similarity_score: int,
    resume_text: str,
    matched_skills: list[str],
) -> list[str]:
    suggestions: list[str] = []
    normalized_resume = normalize_text(resume_text)
    has_metrics = bool(re.search(r"\b\d+%|\b\d+\+|\b\d+\b", resume_text))
    has_projects = "project" in normalized_resume
    has_summary = any(
        marker in normalized_resume
        for marker in ["summary", "profile", "objective", "about"]
    )
    has_experience = any(
        marker in normalized_resume
        for marker in ["experience", "internship", "work history", "employment"]
    )
    has_skills_section = "skills" in normalized_resume
    has_action_verbs = any(verb in normalized_resume for verb in ACTION_VERBS)

    if missing_skills:
        if len(missing_skills) == 1:
            suggestions.append(
                f"The biggest gap is `{missing_skills[0]}`. If you have used it in coursework, projects, or internships, mention it clearly in both the Skills section and one project bullet."
            )
        else:
            suggestions.append(
                "Your resume is missing a few job-critical skills: "
                + ", ".join(missing_skills[:5])
                + ". Add only the ones you genuinely know, and show where you used them."
            )
    else:
        suggestions.append(
            "Your skill match is already solid. The next improvement is to make your experience sound more specific and results-focused."
        )

    if similarity_score < 65:
        suggestions.append(
            "Your wording does not yet sound close enough to the job description. Mirror the role language more naturally in your summary, skills, and project bullets."
        )

    if not has_summary:
        suggestions.append(
            "Add a short 2-3 line professional summary at the top so recruiters understand your profile in the first few seconds."
        )

    if not has_metrics:
        suggestions.append(
            "A human reviewer would want proof of impact. Add numbers like model accuracy, dataset size, response time improvement, or users served."
        )

    if not has_skills_section:
        suggestions.append(
            "Create a clean Skills section with grouped keywords such as languages, ML libraries, databases, and tools. That makes the resume easier for both ATS and recruiters to scan."
        )

    if not has_projects:
        suggestions.append(
            "Add at least one strong project that directly supports this role. For ML roles, mention the problem, dataset, model used, and final result."
        )

    if not has_experience:
        suggestions.append(
            "If you are a fresher, do not leave the experience gap empty. Use internships, academic projects, certifications, hackathons, or research work to show practical exposure."
        )

    if not has_action_verbs:
        suggestions.append(
            "Your bullet points will feel stronger with action verbs such as 'built', 'developed', 'implemented', or 'optimized' instead of plain descriptions."
        )

    if len(matched_skills) < 3:
        suggestions.append(
            "Right now the resume does not show enough overlap with the role. Bring the most relevant tools and concepts closer to the top half of the page."
        )

    suggestions.append(
        "A recruiter should be able to see role fit in 10 seconds. Prioritize your strongest matching skills, best project, and one measurable achievement near the top."
    )

    deduped_suggestions = list(dict.fromkeys(suggestions))
    return deduped_suggestions[:5]


def build_rewrite_examples(
    missing_skills: list[str], matched_skills: list[str], resume_text: str
) -> list[str]:
    examples: list[str] = []
    has_metrics = bool(re.search(r"\b\d+%|\b\d+\+|\b\d+\b", resume_text))
    primary_skill = matched_skills[0] if matched_skills else "machine learning"
    secondary_skill = matched_skills[1] if len(matched_skills) > 1 else "python"

    examples.append(
        "Instead of 'Worked on ML project', write 'Built a "
        + primary_skill
        + " project using "
        + secondary_skill
        + " to solve a real problem and explain the outcome clearly.'"
    )

    if not has_metrics:
        examples.append(
            "Instead of 'Developed a model', write 'Developed a model that achieved measurable results such as accuracy, time reduction, or improved prediction quality.'"
        )

    if missing_skills:
        examples.append(
            "If you know "
            + missing_skills[0]
            + ", mention it with context like 'Used "
            + missing_skills[0]
            + " for data handling, querying, or deployment in an academic or personal project.'"
        )

    examples.append(
        "Keep project bullets in this pattern: action + tool + problem + result. That format sounds much more professional and human."
    )

    return examples[:3]


def calculate_ats_score(similarity_score: int, skill_coverage: int) -> int:
    return round((similarity_score * 0.55) + (skill_coverage * 0.45))


def describe_score(ats_score: int) -> dict[str, str]:
    if ats_score >= 80:
        return {
            "label": "Strong Match",
            "recommendation": "Apply with confidence after a final proofreading pass.",
        }
    if ats_score >= 65:
        return {
            "label": "Good Match",
            "recommendation": "Improve the missing skills and wording before applying.",
        }
    if ats_score >= 45:
        return {
            "label": "Partial Match",
            "recommendation": "Tailor the resume more closely to this job description.",
        }
    return {
        "label": "Needs Tailoring",
        "recommendation": "Revise the resume before applying to this role.",
    }


def build_priority_fixes(
    missing_skills: list[str],
    keyword_overlap: list[str],
    structure_report: dict[str, Any],
    similarity_score: int,
) -> list[str]:
    fixes: list[str] = []
    failed_checks = [
        check for check in structure_report["checks"] if check["status"] != "Passed"
    ]

    if missing_skills:
        fixes.append(
            "Add missing role skills naturally: "
            + ", ".join(missing_skills[:5])
            + ". Mention only skills you can explain in an interview."
        )

    if similarity_score < 65:
        fixes.append(
            "Rewrite the summary and project bullets using language from the job description."
        )

    for check in failed_checks[:2]:
        fixes.append(check["detail"])

    if len(keyword_overlap) < 4:
        fixes.append(
            "Increase keyword overlap by adding relevant tools, responsibilities, and domain words from the job post."
        )

    fixes.append(
        "Keep every change honest. ATS optimization should improve clarity, not add fake experience."
    )
    return list(dict.fromkeys(fixes))[:5]


def analyze_resume_job_match(resume_text: str, job_description: str) -> dict[str, Any]:
    normalized_resume = normalize_text(resume_text)
    normalized_job = normalize_text(job_description)

    resume_skills = extract_skills(normalized_resume)
    job_skills = extract_skills(normalized_job)

    matched_skills = sorted(set(resume_skills) & set(job_skills))
    missing_skills = sorted(set(job_skills) - set(resume_skills))

    similarity_score = compute_similarity_score(normalized_resume, normalized_job)
    skill_coverage = (
        round((len(matched_skills) / len(job_skills)) * 100) if job_skills else 0
    )
    ats_score = calculate_ats_score(similarity_score, skill_coverage)

    resume_keywords = extract_keywords(resume_text)
    job_keywords = extract_keywords(job_description)
    keyword_overlap = sorted(set(resume_keywords) & set(job_keywords))
    missing_keywords = sorted(set(job_keywords) - set(resume_keywords))
    structure_report = analyze_resume_structure(resume_text)
    score_summary = describe_score(ats_score)

    strengths = generate_strengths(similarity_score, matched_skills, resume_text)
    suggestions = generate_suggestions(
        missing_skills, similarity_score, resume_text, matched_skills
    )
    rewrite_examples = build_rewrite_examples(
        missing_skills, matched_skills, resume_text
    )
    priority_fixes = build_priority_fixes(
        missing_skills, keyword_overlap, structure_report, similarity_score
    )

    return {
        "ats_score": ats_score,
        "score_label": score_summary["label"],
        "recommendation": score_summary["recommendation"],
        "similarity_score": similarity_score,
        "skill_coverage": skill_coverage,
        "readability_score": structure_report["readability_score"],
        "word_count": structure_report["word_count"],
        "length_message": structure_report["length_message"],
        "resume_skills": resume_skills,
        "job_skills": job_skills,
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "missing_keywords": missing_keywords[:10],
        "keyword_overlap": keyword_overlap,
        "formatting_checks": structure_report["checks"],
        "section_presence": structure_report["sections"],
        "strengths": strengths,
        "suggestions": suggestions,
        "rewrite_examples": rewrite_examples,
        "priority_fixes": priority_fixes,
    }


def extract_text_from_bytes(filename: str, file_bytes: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(file_bytes))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        except Exception:
            return ""

    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(io.BytesIO(file_bytes))
            return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
        except Exception:
            return ""

    return ""


def extract_text_from_uploaded_file(uploaded_file: Any) -> str:
    filename = getattr(uploaded_file, "filename", None) or getattr(
        uploaded_file, "name", ""
    )
    if hasattr(uploaded_file, "read"):
        file_bytes = uploaded_file.read()
    else:
        file_bytes = uploaded_file.getvalue()
    return extract_text_from_bytes(filename, file_bytes)


def extract_text_from_path(file_path: str | Path) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        except Exception:
            return ""

    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
        except Exception:
            return ""

    return ""
